"""Excel export -- optional.

openpyxl is a guarded import: if it is not installed the dashboard simply
offers CSV, rather than the app failing to start.
"""

from __future__ import annotations

from pathlib import Path

from ..core import aggregate
from ..core.clock import day_range, local_dt
from ..store import repo_sessions
from .csv_export import _target_dir, round_seconds


def available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def write_workbook(engine, start_day: str, end_day: str, *, out=None,
                   rounding_min: int = 0, exclude_unreviewed: bool = False) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    days = day_range(start_day, end_day)
    per_day = aggregate.totals_for_days(engine.db, days)
    combined = aggregate.totals_for_range(engine.db, start_day, end_day)

    wb = Workbook()
    bold = Font(bold=True)

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = f"TimeSplit — {start_day} to {end_day}"
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([])
    summary.append(["Job", "Hours", "Sessions", "Unconfirmed hours"])
    for cell in summary[3]:
        cell.font = bold
    for cat in combined.categories:
        seconds = cat.seconds - (cat.review_seconds if exclude_unreviewed else 0.0)
        summary.append([
            cat.name,
            round(round_seconds(max(0.0, seconds), rounding_min) / 3600.0, 2),
            cat.sessions,
            round(cat.review_seconds / 3600.0, 2),
        ])
    first_data_row = 4
    last_data_row = 3 + len(combined.categories)
    if combined.categories:
        chart = BarChart()
        chart.title = "Hours by job"
        chart.y_axis.title = "Hours"
        chart.add_data(
            Reference(summary, min_col=2, min_row=3, max_row=last_data_row), titles_from_data=True
        )
        chart.set_categories(
            Reference(summary, min_col=1, min_row=first_data_row, max_row=last_data_row)
        )
        chart.height, chart.width = 7, 14
        summary.add_chart(chart, "F3")

    daily = wb.create_sheet("Daily")
    daily.append(["Date", "Job", "Hours", "Sessions", "Unconfirmed hours"])
    for cell in daily[1]:
        cell.font = bold
    for day_total in per_day:
        for cat in day_total.categories:
            daily.append([
                day_total.day, cat.name, round(cat.seconds / 3600.0, 2),
                cat.sessions, round(cat.review_seconds / 3600.0, 2),
            ])

    detail = wb.create_sheet("Detail")
    detail.append([
        "Date", "Start", "End", "Hours", "Job", "App", "Website", "Window title",
        "Assigned by", "Unconfirmed",
    ])
    for cell in detail[1]:
        cell.font = bold
    for row in repo_sessions.sessions_between(engine.db, start_day, end_day):
        if exclude_unreviewed and row["needs_review"]:
            continue
        detail.append([
            row["local_day"],
            local_dt(float(row["start_ts"]), engine.tz).strftime("%H:%M:%S"),
            local_dt(float(row["end_ts"]), engine.tz).strftime("%H:%M:%S"),
            round(float(row["duration_s"]) / 3600.0, 3),
            row.get("category_name") or "Uncategorised",
            row.get("exe_name") or "", row.get("domain") or "",
            row.get("title") or "", row["source"],
            "yes" if row["needs_review"] else "",
        ])

    for sheet, widths in (
        (summary, [26, 10, 11, 18]),
        (daily, [12, 22, 9, 10, 18]),
        (detail, [12, 10, 10, 9, 20, 18, 24, 60, 14, 12]),
    ):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
    summary.freeze_panes = "A4"
    summary["A1"].alignment = Alignment(vertical="center")

    target = _target_dir(out) / f"timesplit_{start_day}_to_{end_day}.xlsx"
    wb.save(target)
    return target
